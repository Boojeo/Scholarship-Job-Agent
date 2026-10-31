"""Markdown -> PDF (xhtml2pdf, pure Python, works on Windows) and -> DOCX (python-docx)."""
from __future__ import annotations

import re
from pathlib import Path

import markdown
from docx import Document
from docx.shared import Pt
from xhtml2pdf import pisa

CSS = """
@page { size: A4; margin: 1.4cm 1.6cm; }
body { font-family: Helvetica; font-size: 10pt; line-height: 1.3; color: #111; }
h1 { font-size: 18pt; margin: 0 0 2pt 0; }
h2 { font-size: 11.5pt; border-bottom: 1px solid #555; margin: 9pt 0 3pt 0; padding-bottom: 1pt; }
p { margin: 2pt 0; } ul { margin: 1pt 0 3pt 14pt; padding: 0; } li { margin: 0 0 1pt 0; }
"""


def md_to_pdf(md_text: str, out: Path) -> Path:
    html = f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>" \
           f"{markdown.markdown(md_text)}</body></html>"
    with open(out, "wb") as fh:
        result = pisa.CreatePDF(html, dest=fh, encoding="utf-8")
    if result.err:
        raise RuntimeError(f"PDF rendering failed for {out.name}")
    return out


def _add_runs(par, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            par.add_run(part).bold = bool(i % 2)


def md_to_docx(md_text: str, out: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\s*[-*] ", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        else:
            _add_runs(doc.add_paragraph(), line)
    doc.save(out)
    return out


def render_both(md_text: str, base: Path) -> tuple[Path, Path]:
    base.parent.mkdir(parents=True, exist_ok=True)
    return md_to_pdf(md_text, base.with_suffix(".pdf")), md_to_docx(md_text, base.with_suffix(".docx"))
